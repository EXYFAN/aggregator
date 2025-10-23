
import base64
import math
import os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_docx(text_file, output_filename):
    """
    Generates a .docx file from a text file with specific formatting.
    """
    document = Document()

    # Set page margins to 2.5cm
    sections = document.sections
    for section in sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Set default font to SimSun
    style = document.styles['Normal']
    style.font.name = 'SimSun'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    with open(text_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Add book title
    book_title = lines[0].strip()
    p = document.add_paragraph()
    runner = p.add_run(book_title)
    runner.font.name = 'SimSun'
    runner.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    runner.bold = True
    runner.font.size = Pt(12)

    # Add chapter title
    chapter_title = lines[2].strip()
    h1 = document.add_heading(chapter_title, level=1)
    h1.runs[0].font.name = 'SimSun'
    h1.runs[0].element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    h1.runs[0].bold = True
    h1.runs[0].font.size = Pt(16)
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Add body text
    for line in lines[4:]:
        line = line.strip()
        if line:
            p = document.add_paragraph()
            run = p.add_run(line)
            run.font.name = 'SimSun'
            p.style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(12)

    # Add chapter end
    p = document.add_paragraph()
    run = p.add_run('—— 第一章完')
    run.font.name = 'SimSun'
    p.style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(12)

    # Add footer with page numbers
    section = document.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar1)

    run2 = p.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    run2._r.append(instrText)

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run2._r.append(fldChar3)
    
    document.save(output_filename)

def encode_and_chunk(filename, chunk_size=150000):
    """
    Base64-encodes a file and splits it into chunks.
    """
    with open(filename, 'rb') as f:
        encoded_string = base64.b64encode(f.read()).decode('utf-8')

    num_chunks = math.ceil(len(encoded_string) / chunk_size)

    for i in range(num_chunks):
        start = i * chunk_size
        end = start + chunk_size
        chunk = encoded_string[start:end]
        print(f"Part {i+1}/{num_chunks}:")
        print(chunk)

if __name__ == "__main__":
    story_file = os.path.join(os.path.dirname(__file__), 'story.txt')
    docx_file = '校准修仙-第一章.docx'
    create_docx(story_file, docx_file)
    encode_and_chunk(docx_file)
    os.remove(docx_file)
